"""Daily price table (data/daily.csv) plus the files the website and users download."""
import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from products import PRODUCTS
from sources import Price

DATA = Path(__file__).resolve().parent.parent / "data"
DAILY = DATA / "daily.csv"
FIELDS = ["date", "brand", "product", "name", "price"]
BRANDS = {"ptt": "ปตท. (PTT)", "bangchak": "บางจาก"}

Key = Tuple[str, str, str]  # (date, brand, product)


def load() -> Dict[Key, dict]:
    if not DAILY.exists():
        return {}
    with DAILY.open(encoding="utf-8") as f:
        return {(r["date"], r["brand"], r["product"]): r for r in csv.DictReader(f)}


def upsert(rows: Dict[Key, dict], day: date, prices: List[Price]) -> int:
    """Set the prices in effect on `day`. Returns how many rows were new or changed."""
    changed = 0
    for p in prices:
        key = (day.isoformat(), p.brand, p.product)
        row = {"date": key[0], "brand": p.brand, "product": p.product, "name": p.name, "price": f"{p.price:.2f}"}
        if rows.get(key) != row:
            rows[key] = row
            changed += 1
    return changed


def save(rows: Dict[Key, dict]) -> None:
    order = list(PRODUCTS)
    rank = lambda r: (r["date"], r["brand"], order.index(r["product"]) if r["product"] in order else 99, r["product"])
    DATA.mkdir(exist_ok=True)
    with DAILY.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(sorted(rows.values(), key=rank))


def write_latest(prices: List[Price], updated: datetime, day: date) -> dict:
    """data/latest.json: what the website shows (today's prices, tomorrow's if announced)."""
    latest = {
        "updated": updated.isoformat(timespec="minutes"),
        "date": day.isoformat(),
        "brands": {b: {"name": label, "products": []} for b, label in BRANDS.items()},
    }
    for p in prices:
        latest["brands"][p.brand]["products"].append({
            "product": p.product,
            "label": PRODUCTS.get(p.product, p.name),
            "name": p.name,
            "price": p.price,
            "since": p.effective.isoformat(timespec="minutes") if p.effective else None,
            "tomorrow": p.tomorrow,
        })
    (DATA / "latest.json").write_text(json.dumps(latest, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return latest


def write_excel(rows: Dict[Key, dict], latest: dict) -> Path:
    """data/thai-oil-prices.xlsx: today's comparison, a daily sheet per brand and the raw table."""
    df = pd.DataFrame(rows.values())
    df["price"] = df["price"].astype(float)
    df["น้ำมัน"] = df["product"].map(lambda k: PRODUCTS.get(k, k))
    order = [PRODUCTS[k] for k in PRODUCTS]

    today = {}
    for brand, info in latest["brands"].items():
        for p in info["products"]:
            today.setdefault(p["label"], {})[BRANDS[brand]] = p["price"]
            if p["tomorrow"] is not None:
                today[p["label"]][f"{BRANDS[brand]} พรุ่งนี้"] = p["tomorrow"]
    today_df = pd.DataFrame.from_dict(today, orient="index")
    today_df = today_df.reindex([n for n in order if n in today_df.index])
    today_df.index.name = "น้ำมัน"

    out = DATA / "thai-oil-prices.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        today_df.to_excel(xl, sheet_name="วันนี้")
        for brand, label in BRANDS.items():
            part = df[df["brand"] == brand]
            if part.empty:
                continue
            wide = part.pivot(index="date", columns="น้ำมัน", values="price")
            wide = wide[[c for c in order if c in wide.columns]].sort_index(ascending=False)
            wide.index.name = "วันที่"
            wide.to_excel(xl, sheet_name=f"{label} รายวัน"[:31])
        raw = df[["date", "brand", "น้ำมัน", "name", "price"]].rename(columns={
            "date": "วันที่", "brand": "แบรนด์", "name": "ชื่อที่แบรนด์ใช้", "price": "ราคา (บาท/ลิตร)"})
        raw.sort_values(["วันที่", "แบรนด์"], ascending=[False, True]).to_excel(xl, sheet_name="ข้อมูลทั้งหมด", index=False)
        for ws in xl.book.worksheets:
            _style(ws)
    return out


def _style(ws) -> None:
    head = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(10, width + 2))
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = "0.00"
